"""Letter generation.

Accepted argument nodes in, a letter out, with **every paragraph traceable to an
argument node or to a fixed procedural template**.

The §4.4 check, and how it reads here
------------------------------------
The spec says every paragraph must carry an ``argument_node_id`` and that any
paragraph without one is dropped or flagged. Taken literally that deletes the
letter's procedural header — the appeal-rights statement, the deadline
reference, the claim identifiers — which comes from no argument and is legally
necessary.

The design anticipated this: `components.md` gives the `Paragraph` component a
badge reading ``procedural`` "for the ones no argument produced". So the check
implemented here is:

    a paragraph is valid iff it carries an argument_node_id OR the reserved
    `procedural` tag, and `procedural` paragraphs are rendered from a fixed
    template, never from model prose.

LLM-generated text carrying neither tag is dropped. That keeps the spec's
guarantee — no unattributable factual claim reaches the letter — without
deleting the parts of the letter that make it a valid filing.
"""

from __future__ import annotations

import logging
from dataclasses import dataclass
from datetime import UTC, date, datetime
from typing import Any, Literal

from app.db import repo
from app.domain.argument import ArgumentGraph
from app.domain.case import CaseFacts, CaseStage
from app.domain.route import RouteDetermination
from app.problem import ErrorCode, Problem
from app.services import llm, storage
from app.site import DISCLAIMER

logger = logging.getLogger("appeal_architect.letters")

#: The reserved tag for a paragraph no argument produced.
PROCEDURAL = "procedural"

ParagraphSource = Literal["argument", "procedural"]


@dataclass
class Paragraph:
    """One paragraph, and what produced it."""

    text: str
    source: ParagraphSource
    argument_node_id: str | None = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "text": self.text,
            "source": self.source,
            "argument_node_id": self.argument_node_id or PROCEDURAL,
        }


def _opening(facts: CaseFacts, route: RouteDetermination) -> list[Paragraph]:
    """The procedural header. Fixed template, never model prose.

    It states what the letter is, identifies the claim, and cites the appeal
    right being exercised. A reviewer looks for exactly this.
    """
    claim_line = (
        f"Claim number {facts.claim_number}."
        if facts.claim_number
        else "Claim number: see enclosed."
    )
    internal = next((d for d in route.deadlines if d.id == "internal_appeal"), None)

    paragraphs = [
        Paragraph(
            text=(
                f"I am writing to appeal {facts.insurer_name}'s denial of coverage dated "
                f"{facts.denial_date:%-d %B %Y}. {claim_line} I am asking you to review "
                "that decision and to overturn it for the reasons set out below."
            ),
            source="procedural",
        )
    ]

    if internal is not None:
        rights = (
            f"This appeal is made under {internal.citation.label()}, which provides the "
            f"right to appeal an adverse benefit determination. By my calculation the "
            f"window for this appeal runs to {internal.due_date:%-d %B %Y}."
        )
        if internal.ambiguous:
            rights += (
                " That date is counted from the date printed on your letter, because I "
                "do not have a record of the date it reached me; if you calculate it "
                "differently, please tell me in writing."
            )
        paragraphs.append(Paragraph(text=rights, source="procedural"))

    paragraphs.append(
        Paragraph(
            text=(
                "Under 29 C.F.R. § 2560.503-1(h), I request free copies of all documents, "
                "records and other information relied on in making this determination, "
                "including any internal clinical criteria or guidelines applied to my claim."
            ),
            source="procedural",
        )
    )
    return paragraphs


def _closing(facts: CaseFacts, route: RouteDetermination) -> list[Paragraph]:
    response = next((d for d in route.deadlines if d.id == "insurer_response"), None)
    ask = (
        "Please confirm in writing that you have received this appeal, and send your "
        "decision in writing."
    )
    if response is not None:
        ask += (
            f" I understand you are required to respond by {response.due_date:%-d %B %Y} "
            f"under {response.citation.label()}."
        )
    return [
        Paragraph(text=ask, source="procedural"),
        Paragraph(
            text=(
                "If you uphold this denial, please include a statement of my right to "
                "external review and the steps for requesting it."
            ),
            source="procedural",
        ),
    ]


def _accepted_arguments(graph: ArgumentGraph) -> list[dict[str, Any]]:
    """Solid ground first, then worth adding. Defeated arguments are not sent.

    The order is the argument: a reviewer reads the strongest point first, and
    an argument the solver found contestable belongs after the ones that hold.
    """
    ordered: list[dict[str, Any]] = []
    for node_id in list(graph.grounded_extension) + list(graph.worth_adding):
        argument = graph.by_id(node_id)
        if argument is None:  # pragma: no cover
            continue
        ordered.append(
            {
                "argument_node_id": argument.id,
                "tier": graph.tier_of(argument.id),
                "claim": argument.claim,
                "premises": [
                    {"text": p.text, "evidence_on_file": p.satisfied} for p in argument.premises
                ],
                "citations": [c.label() for c in argument.citations],
            }
        )
    return ordered


def check_paragraphs(paragraphs: list[Paragraph]) -> list[Paragraph]:
    """Drop any paragraph that is neither tagged to an argument nor procedural.

    This is the §4.4 gate. A paragraph of model prose we cannot attribute to an
    argument node is exactly the thing that would put an unsupported factual
    claim into a legal filing, so it is dropped rather than flagged and shipped.
    """
    kept: list[Paragraph] = []
    for paragraph in paragraphs:
        if paragraph.source == "procedural":
            kept.append(paragraph)
            continue
        if paragraph.argument_node_id:
            kept.append(paragraph)
            continue
        logger.warning("dropped an untagged generated paragraph")
    return kept


def generate_letter(case_id: str, user_id: str) -> dict[str, Any]:
    """Build the next version of the appeal letter."""
    from app.services import arguments as arguments_service
    from app.services import facts as facts_service
    from app.services import route as route_service

    facts = facts_service.build_case_facts(case_id, user_id)

    stored_route = route_service.get_stored(case_id, user_id)
    route = (
        route_service._from_row(stored_route)
        if stored_route
        else route_service.compute_route(case_id, user_id)
    )

    stored_graph = arguments_service.get_stored(case_id, user_id)
    if stored_graph is None:
        graph = arguments_service.compute_graph(case_id, user_id)
    else:
        from app.api.v1.arguments import _from_row

        graph = _from_row(stored_graph)

    accepted = _accepted_arguments(graph)
    if not accepted:
        raise Problem(
            409,
            ErrorCode.CONFLICT,
            "There is no argument to write from yet. Confirm the facts and let us work "
            "out the arguments first.",
        )

    llm.check_monthly_cap(user_id)
    generated, _ = llm.write_letter_paragraphs(
        arguments=accepted,
        confirmed_facts={
            "insurer_name": facts.insurer_name,
            "claim_number": facts.claim_number,
            "denial_date": facts.denial_date.isoformat(),
            "service_date": facts.service_date.isoformat() if facts.service_date else None,
            "state": facts.state,
            "denial_reasons": [r.value for r in facts.denial_reasons],
            "cited_policy_language": facts.cited_policy_language,
        },
        user_id=user_id,
    )

    known_ids = {item["argument_node_id"] for item in accepted}
    body = [
        Paragraph(
            text=item["text"],
            source="argument",
            argument_node_id=item["argument_node_id"],
        )
        for item in generated
        # A node id the model invented is not an argument the solver accepted, so
        # the paragraph has nothing behind it.
        if item["argument_node_id"] in known_ids
    ]

    paragraphs = check_paragraphs([*_opening(facts, route), *body, *_closing(facts, route)])

    existing = repo.list_for_case("letters", case_id, user_id, order="version", desc=True)
    version = (existing[0]["version"] + 1) if existing else 1

    row = repo.insert_for_case(
        "letters",
        case_id,
        user_id,
        [
            {
                "argument_graph_id": stored_graph["id"] if stored_graph else None,
                "version": version,
                "body": {
                    "paragraphs": [p.to_dict() for p in paragraphs],
                    "disclaimer": DISCLAIMER,
                    "rulebase_version": route.rulebase_version,
                    "schemes_version": graph.schemes_version,
                    "generated_for": facts.insurer_name,
                },
            }
        ],
    )[0]

    repo.update_case(case_id, user_id, {"stage": CaseStage.LETTER_READY.value})
    repo.log_event(
        case_id,
        "letter_generated",
        {
            "version": version,
            "paragraphs": len(paragraphs),
            "argument_paragraphs": sum(1 for p in paragraphs if p.source == "argument"),
        },
    )
    return row


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------


def render_html(letter: dict[str, Any], case: dict[str, Any]) -> str:
    """The letter as printable HTML, US Letter with 1 inch margins.

    Node badges are deliberately absent: they are an editing affordance, not
    part of the document that gets filed.
    """
    body = letter.get("body") or {}
    paragraphs = body.get("paragraphs") or []
    today = date.today()

    def escape(value: object) -> str:
        return str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    rows = "\n".join(f'  <p class="body">{escape(p.get("text", ""))}</p>' for p in paragraphs)

    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<title>Appeal letter</title>
<style>
  @page {{ size: Letter; margin: 1in; }}
  body {{ font-family: Georgia, "Times New Roman", serif; font-size: 11.5pt;
          line-height: 1.55; color: #1C1A17; }}
  .meta {{ font-family: "Courier New", monospace; font-size: 9.5pt;
           color: #5E5850; margin-bottom: 2em; }}
  .meta div {{ margin: 0.15em 0; }}
  p.body {{ margin: 0 0 1em; text-align: left; }}
  .sign {{ margin-top: 2.5em; }}
  .disclaimer {{ margin-top: 3em; padding-top: 0.8em;
                 border-top: 1px solid #DED5C6; font-size: 9pt; color: #1C1A17; }}
</style></head>
<body>
  <div class="meta">
    <div>{escape(today.strftime("%-d %B %Y"))}</div>
    <div>{escape(case.get("insurer_name") or "")}</div>
    <div>Appeals Department</div>
    <div>Claim: {escape(case.get("claim_number") or "see enclosed")}</div>
  </div>

  <p class="body">To the Appeals Department,</p>
{rows}
  <div class="sign">
    <p class="body">Sincerely,</p>
    <p class="body">&nbsp;</p>
    <p class="body">_______________________________</p>
    <p class="body">Member signature and date</p>
  </div>
  <div class="disclaimer">{escape(DISCLAIMER)}</div>
</body></html>
"""


def render_pdf(html: str) -> bytes:
    from weasyprint import HTML

    return bytes(HTML(string=html).write_pdf() or b"")


def render_docx(letter: dict[str, Any], case: dict[str, Any]) -> bytes:
    """DOCX, so the user can edit it in Word before filing."""
    import io

    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        for attribute in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, attribute, Inches(1))

    style = document.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(11.5)

    header = document.add_paragraph()
    header.add_run(date.today().strftime("%-d %B %Y")).italic = True
    document.add_paragraph(str(case.get("insurer_name") or ""))
    document.add_paragraph("Appeals Department")
    document.add_paragraph(f"Claim: {case.get('claim_number') or 'see enclosed'}")
    document.add_paragraph("")
    document.add_paragraph("To the Appeals Department,")

    for paragraph in (letter.get("body") or {}).get("paragraphs") or []:
        document.add_paragraph(str(paragraph.get("text", "")))

    document.add_paragraph("")
    document.add_paragraph("Sincerely,")
    document.add_paragraph("")
    document.add_paragraph("_______________________________")
    document.add_paragraph("Member signature and date")
    document.add_paragraph("")
    disclaimer = document.add_paragraph()
    run = disclaimer.add_run(DISCLAIMER)
    run.font.size = Pt(9)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export(case_id: str, letter_id: str, user_id: str, fmt: str) -> str:
    """Render, store and return a signed URL for the chosen format."""
    # Validated before any I/O: a bad format is a bad request, and answering it
    # should not cost a database round trip or depend on the database being up.
    if fmt not in {"pdf", "docx"}:
        raise Problem(422, ErrorCode.VALIDATION_FAILED, "Export as pdf or docx.")

    letter = repo.get_child("letters", letter_id, case_id, user_id)
    case = repo.get_case(case_id, user_id)

    if fmt == "pdf":
        data = render_pdf(render_html(letter, case))
        content_type = "application/pdf"
        column = "storage_path_pdf"
        suffix = "pdf"
    else:
        data = render_docx(letter, case)
        content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        column = "storage_path_docx"
        suffix = "docx"

    path = f"{case_id}/appeal-v{letter['version']}-{datetime.now(tz=UTC):%Y%m%d%H%M%S}.{suffix}"
    storage.upload("letters", path, data, content_type)
    repo.update_child("letters", letter_id, case_id, user_id, {column: path})
    repo.log_event(case_id, "letter_exported", {"format": fmt, "version": letter["version"]})
    return storage.signed_url("letters", path)
